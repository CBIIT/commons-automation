"""
Execute Steps sheet rows via the app-specific ExcelCaseProfile.

Shared actions live in _common_steps; app-only actions in profile.extra_handlers.
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

import openpyxl

from Commons.excel_testcase._common_steps import common_handlers, norm_action
from Commons.excel_testcase.registry import get_profile_for_app, profile_for_case_path

if TYPE_CHECKING:
    from Commons.apps.base import BaseApp

_COL_INPUT = 0
_COL_ACTION = 1
_COL_PAGE = 2
_COL_PARAMS = 3
_COL_FILTER = 4


def _input_excel_from_config(cfg: tuple) -> str:
    if not cfg:
        return ""
    if len(cfg) > _COL_INPUT and cfg[_COL_INPUT]:
        return str(cfg[_COL_INPUT]).strip()
    if len(cfg) > 2 and cfg[2]:
        return str(cfg[2]).strip()
    return ""


def _dispatch_step(app: "BaseApp", profile, excel_name: str, action: str, page: str, params: str, flt: str) -> None:
    common = common_handlers()
    if action in profile.extra_handlers:
        profile.extra_handlers[action](app, page, params, flt, excel_name)
        return
    if action in common:
        common[action](app, page, params, flt, excel_name)
        return
    raise ValueError(
        f"Unknown step action {action!r} in {excel_name} for {profile.app_key}. "
        f"Add a handler in Commons/excel_testcase/apps/{profile.testcases_dir.lower()}.py "
        f"or Commons/excel_testcase/_common_steps.py"
    )


def run_testcase_excel(app: "BaseApp", excel_path: Path) -> None:
    """Run one study-folder TC*.xlsx using that app's ExcelCaseProfile."""
    if not excel_path.is_file():
        raise FileNotFoundError(excel_path)

    wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    if "Steps" not in wb.sheetnames:
        wb.close()
        raise KeyError(f"{excel_path.name} must contain a 'Steps' sheet")
    ws = wb["Steps"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        raise ValueError(f"{excel_path.name}: Steps sheet is empty")

    profile = profile_for_case_path(excel_path) or get_profile_for_app(app)

    input_excel = _input_excel_from_config(rows[1])
    if input_excel:
        input_name = Path(input_excel.replace("\\", "/")).name
        # Static page workbooks (*_StaticData.xlsx) are not TabQuery TC*.xlsx files.
        if not input_name.endswith("_StaticData.xlsx"):
            app.load_query_file(input_name)

    for row in rows[2:]:
        if not row or len(row) <= _COL_ACTION or not row[_COL_ACTION]:
            continue
        action = norm_action(str(row[_COL_ACTION]))
        page = str(row[_COL_PAGE]).strip() if len(row) > _COL_PAGE and row[_COL_PAGE] else ""
        params = str(row[_COL_PARAMS]).strip() if len(row) > _COL_PARAMS and row[_COL_PARAMS] else ""
        flt = str(row[_COL_FILTER]).strip() if len(row) > _COL_FILTER and row[_COL_FILTER] else ""
        _dispatch_step(app, profile, excel_path.name, action, page, params, flt)
