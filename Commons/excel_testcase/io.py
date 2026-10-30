"""Write TC*.xlsx testcase workbooks (Steps sheet + per-app Action dropdown)."""

from __future__ import annotations

from pathlib import Path
from typing import Sequence

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from Commons.excel_testcase.profile import StepRow

STEPS_HEADER = ["Input Excel", "Action", "Page", "Params", "Filter"]
_HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
_DROPDOWN_ROWS = 50


def write_excel_case(
    out_path: Path,
    *,
    input_excel: str,
    steps: Sequence[StepRow],
    excel_actions: Sequence[str],
    dropdown_name: str = "ExcelCaseActions",
) -> None:
    """Row 2 = query workbook path; following rows = steps. Action list is per-app."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Steps"
    ws.append(list(STEPS_HEADER))
    ws.append([input_excel, None, None, None, None])
    for action, page, params, flt in steps:
        ws.append([None, action, page, params, flt])

    lists = wb.create_sheet("_Lists")
    lists.sheet_state = "hidden"
    for i, act in enumerate(excel_actions, start=1):
        lists.cell(row=i, column=1, value=act)
    end = min(len(excel_actions), _DROPDOWN_ROWS)
    wb.defined_names.add(
        DefinedName(dropdown_name, attr_text=f"_Lists!$A$1:$A${max(end, 1)}")
    )

    dv = DataValidation(type="list", formula1=f"={dropdown_name}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B3:B500")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    wb.save(out_path)
