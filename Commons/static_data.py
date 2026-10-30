"""
Static marketing/info pages: compare UI text and links to InputFiles/<app>/*_StaticData.xlsx.
Each sheet (e.g. HomePage) lists object names, expected copy, and optional link URLs (QA host in file;
stage/prod URLs are derived when the object name contains Link).
"""
import re
from pathlib import Path
from typing import List, TYPE_CHECKING
from urllib.parse import urlparse

import openpyxl

from Commons import action_log, excel_writer
from Commons.object_repository import LocatorNotFound
from Commons.settings import input_dir

if TYPE_CHECKING:
    from Commons.apps.base import BaseApp


# ---------------------------------------------------------------------------
# Text / URL normalization
# ---------------------------------------------------------------------------


def _normalize_text(text: str) -> str:
    """Collapse newlines, non-breaking spaces, and curly quotes (UI inner_text vs Excel)."""
    text = text or ""
    text = text.replace("\u00a0", " ").replace("\u202f", " ").replace("\u200b", "")
    text = text.replace("\u2019", "'").replace("\u2018", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = re.sub(r"\r?\n", " ", text)
    text = re.sub(r" +", " ", text)
    return text.strip()


def _normalize_href(href: str) -> str:
    return (href or "").replace("%20", " ").strip()


def _hash_route(href: str) -> str:
    """'#/support' from either '#/support' or 'https://host/#/support'."""
    if "#/" in href:
        return "#/" + href.split("#/", 1)[1]
    return ""


def _is_site_path(href: str) -> bool:
    """True for same-site paths like /about (not //cdn or a full URL)."""
    return href.startswith("/") and not href.startswith("//")


def _url_path(href: str) -> str:
    parsed = urlparse(href)
    path = parsed.path or "/"
    if parsed.query:
        path = f"{path}?{parsed.query}"
    return path.rstrip("/") or "/"


def _path_matches_absolute(maybe_path: str, maybe_abs: str) -> bool:
    """'/about' matches 'https://host/about'. Do not treat '/' as every homepage."""
    if not _is_site_path(maybe_path) or not maybe_abs.startswith("http"):
        return False
    path = _url_path(maybe_path)
    if path == "/":
        return False
    return path == _url_path(maybe_abs)


def _hrefs_match(got: str, want: str) -> bool:
    """True if hrefs are the same after trim, trailing slash, hash routes, or path vs full URL."""
    got = _normalize_href(got)
    want = _normalize_href(want)
    if got == want:
        return True
    if got.rstrip("/") == want.rstrip("/"):
        return True
    got_hash, want_hash = _hash_route(got), _hash_route(want)
    if got_hash and want_hash and got_hash == want_hash:
        return True
    if _path_matches_absolute(got, want) or _path_matches_absolute(want, got):
        return True
    return False


def _link_href_for_env(qa_href: str, env: str) -> str:
    """Map QA link from the sheet to the current test environment."""
    if env == "qa":
        return qa_href
    if env == "stage":
        return qa_href.replace("-qa", "-stage")
    return qa_href.replace("-qa", "")  # prod


# ---------------------------------------------------------------------------
# Static page verification (called from BaseApp.verify_static_data)
# ---------------------------------------------------------------------------


def _static_log(sheet_name: str, obj: str, msg: str) -> None:
    action_log.static_log(f"[Static {sheet_name}] {obj}: {msg}")


def verify_static_page(app: "BaseApp", sheet_name: str) -> List[str]:
    """Check one Excel sheet against Static/<sheet_name>/ locators; return failure messages (empty = pass)."""
    path = Path(input_dir(app.NAME)) / app.STATIC_DATA_EXCEL
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Sheet '{sheet_name}' not in {path.name}. Sheets: {wb.sheetnames}")
    rows = list(wb[sheet_name].iter_rows(values_only=True))[1:]
    wb.close()

    objects = [str(r[0]).strip() for r in rows if r and r[0]]
    action_log.static_log(f"[Static] Verifying sheet {sheet_name!r} ({len(objects)} object(s))")

    problems: List[str] = []
    out_rows: List[list] = []
    for row in rows:
        if not row or not row[0]:
            continue
        obj = str(row[0]).strip()
        expected_text = _normalize_text(str(row[1]) if len(row) > 1 and row[1] is not None else "")
        expected_url = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        ok = False
        try:
            xp = app.repo.xpath(f"Static/{sheet_name}/{obj}")
            el = app.ui.loc(xp).first
            el.wait_for(state="visible", timeout=10000)
            actual = _normalize_text(el.inner_text())
            text_ok = actual == expected_text
            status = "PASS" if text_ok else "FAIL"
            _static_log(
                sheet_name,
                obj,
                f"text {status} — actual={actual!r} expected={expected_text!r}",
            )
            if not text_ok:
                problems.append(f"{obj}: text is '{actual}', expected '{expected_text}'")
            url_ok = True
            if "Link" in obj:
                want = _link_href_for_env(expected_url, app.settings.env)
                got = el.get_attribute("href") or ""
                url_ok = _hrefs_match(got, want)
                status = "PASS" if url_ok else "FAIL"
                _static_log(
                    sheet_name,
                    obj,
                    f"link {status} — actual={got!r} expected={want!r}",
                )
                if not url_ok:
                    problems.append(f"{obj}: link is '{got}', expected '{want}'")
            ok = text_ok and url_ok
        except LocatorNotFound as e:
            _static_log(sheet_name, obj, f"FAIL — {e}")
            problems.append(f"{obj}: {e}")
        except Exception as e:
            _static_log(sheet_name, obj, f"FAIL — element not found or not visible ({type(e).__name__})")
            problems.append(f"{obj}: element not found or not visible ({type(e).__name__})")
        out_rows.append([obj, expected_text, expected_url, str(ok).lower()])

    excel_writer.write_sheet(
        app.output_dir / f"Output_{app.STATIC_DATA_EXCEL}",
        sheet_name,
        ["Object", "Expected text", "Expected URL", "Passed"],
        out_rows,
    )
    return problems
