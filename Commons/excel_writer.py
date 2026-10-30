"""Write UI/DB compare tables to Excel under OutputFiles/<app>/<env>/ (gitignored)."""
from pathlib import Path
from typing import List, Optional, Sequence

import openpyxl
from openpyxl.styles import Font, PatternFill

# ---------------------------------------------------------------------------
# Sheet naming and styling
# ---------------------------------------------------------------------------

_ILLEGAL_SHEET_CHARS = "[]:*?/\\"
LISTING_HEADER_COLOR = "9BC2E6"

# ---------------------------------------------------------------------------
# Write workbook
# ---------------------------------------------------------------------------


def write_sheet(
    path: Path,
    sheet: str,
    header: Sequence[str],
    rows: List[Sequence[str]],
    header_color: Optional[str] = None,
) -> None:
    """Create or replace one worksheet (header + rows) in an .xlsx file."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    sheet = "".join("_" if c in _ILLEGAL_SHEET_CHARS else c for c in sheet)[:31]
    if path.exists():
        wb = openpyxl.load_workbook(path)
        if sheet in wb.sheetnames:
            del wb[sheet]
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    ws = wb.create_sheet(sheet)
    ws.append(list(header))
    for r in rows:
        ws.append([_safe(c) for c in r])
    if header_color:
        fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        font = Font(bold=True, color="1F4E79")
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
    wb.save(path)


# ---------------------------------------------------------------------------
# Cell values
# ---------------------------------------------------------------------------


def _safe(value):
    """openpyxl rejects control characters that innerText occasionally contains."""
    if isinstance(value, str):
        return "".join(ch for ch in value if ch == "\n" or ch == "\t" or ord(ch) >= 32)
    return value
