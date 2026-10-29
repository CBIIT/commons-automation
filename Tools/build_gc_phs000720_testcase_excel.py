#!/usr/bin/env python3
"""
Build one *.xlsx per phs000720 case in Testcases/GC/studies/phs000720/ (next to test_*.py).

Non-technical editors use the Action dropdown on the Steps sheet; filter steps use
Page = filter group, Filter = checkbox value (same as select_filter in Python).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Testcases" / "GC" / "studies" / "phs000720"

HEADER = ["Input Excel", "Action", "Page", "Params", "Filter"]

ACTIONS = [
    "accept_banner",
    "open_data",
    "select_filter",
    "read_statbar",
    "verify_tab",
    "verify_all_tabs",
    "assert_no_failures",
    "click_and_wait",
    "js_click_and_wait",
]

CASES = [
    {
        "stem": "TC01_GC_phs000720_SmpTumSta-Tumor_SampleType_Analyte",
        "input": "InputFiles/GC/TC01_GC_phs000720_SmpTumSta-Tumor_SampleType_Analyte.xlsx",
        "steps": [
            ('accept_banner', '', '', ''),
            ('open_data', '', '', ''),
            ('select_filter', 'StudyFacet/PHS_Accession', '', 'phs000720'),
            ('select_filter', 'SamplesFacet/SampleTumorStatus', '', 'Tumor'),
            ('select_filter', 'SamplesFacet/SampleType', '', 'Analyte'),
            ('read_statbar', '', '', ''),
            ('verify_tab', '', 'Participants', ''),
            ('verify_tab', '', 'Samples', ''),
            ('verify_tab', '', 'Files', ''),
            ('assert_no_failures', '', '', ''),
        ],
    },
    {
        "stem": "TC02_GC_phs000720_Sex_Male",
        "input": "InputFiles/GC/TC02_GC_phs000720_Sex_Male.xlsx",
        "steps": [
            ('accept_banner', '', '', ''),
            ('open_data', '', '', ''),
            ('select_filter', 'StudyFacet/PHS_Accession', '', 'phs000720'),
            ('select_filter', 'DemographicsFacet/Sex', '', 'Male'),
            ('read_statbar', '', '', ''),
            ('verify_tab', '', 'Participants', ''),
            ('verify_tab', '', 'Samples', ''),
            ('verify_tab', '', 'Files', ''),
            ('assert_no_failures', '', '', ''),
        ],
    },
    {
        "stem": "TC03_GC_phs000720_FileType_TIFF",
        "input": "InputFiles/GC/TC03_GC_phs000720_FileType_TIFF.xlsx",
        "steps": [
            ('accept_banner', '', '', ''),
            ('open_data', '', '', ''),
            ('select_filter', 'StudyFacet/PHS_Accession', '', 'phs000720'),
            ('select_filter', 'FilesFacet/FileType', '', 'tiff'),
            ('read_statbar', '', '', ''),
            ('verify_tab', '', 'Participants', ''),
            ('verify_tab', '', 'Samples', ''),
            ('verify_tab', '', 'Files', ''),
            ('assert_no_failures', '', '', ''),
        ],
    },
    {
        "stem": "TC04_GC_phs000720_SampleTumorStatus_Tumor",
        "input": "InputFiles/GC/TC04_GC_phs000720_SampleTumorStatus_Tumor.xlsx",
        "steps": [
            ('accept_banner', '', '', ''),
            ('open_data', '', '', ''),
            ('select_filter', 'StudyFacet/PHS_Accession', '', 'phs000720'),
            ('select_filter', 'SamplesFacet/SampleTumorStatus', '', 'Tumor'),
            ('read_statbar', '', '', ''),
            ('verify_tab', '', 'Participants', ''),
            ('verify_tab', '', 'Samples', ''),
            ('verify_tab', '', 'Files', ''),
            ('assert_no_failures', '', '', ''),
        ],
    },
    {
        "stem": "TC05_GC_phs000720_SampleType_Analyte",
        "input": "InputFiles/GC/TC05_GC_phs000720_SampleType_Analyte.xlsx",
        "steps": [
            ('accept_banner', '', '', ''),
            ('open_data', '', '', ''),
            ('select_filter', 'StudyFacet/PHS_Accession', '', 'phs000720'),
            ('select_filter', 'SamplesFacet/SampleType', '', 'Analyte'),
            ('read_statbar', '', '', ''),
            ('verify_tab', '', 'Participants', ''),
            ('verify_tab', '', 'Samples', ''),
            ('verify_tab', '', 'Files', ''),
            ('assert_no_failures', '', '', ''),
        ],
    },
    {
        "stem": "TC06_GC_phs000720_PrimDiag_NotApplicable",
        "input": "InputFiles/GC/TC06_GC_phs000720_PrimDiag_NotApplicable.xlsx",
        "steps": [
            ('accept_banner', '', '', ''),
            ('open_data', '', '', ''),
            ('select_filter', 'StudyFacet/PHS_Accession', '', 'phs000720'),
            ('select_filter', 'DiagnosisFacet/PrimaryDiagnosis', '', 'NotApplicable'),
            ('read_statbar', '', '', ''),
            ('verify_tab', '', 'Participants', ''),
            ('verify_tab', '', 'Samples', ''),
            ('verify_tab', '', 'Files', ''),
            ('assert_no_failures', '', '', ''),
        ],
    },
]


def _add_lists_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("_Lists")
    ws.sheet_state = "hidden"
    for i, act in enumerate(ACTIONS, start=1):
        ws.cell(row=i, column=1, value=act)
    wb.defined_names.add(DefinedName("GCExcelActions", attr_text="_Lists!$A$1:$A$50"))


def _write_one(case: dict) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"{case['stem']}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Steps"
    ws.append(HEADER)
    ws.append([case["input"], None, None, None, None])
    for action, page, params, flt in case["steps"]:
        ws.append([None, action, page, params, flt])

    _add_lists_sheet(wb)
    dv = DataValidation(type="list", formula1="=GCExcelActions", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B3:B120")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    wb.save(path)
    print(f"Wrote {path}")


def main() -> None:
    for case in CASES:
        _write_one(case)


if __name__ == "__main__":
    main()
