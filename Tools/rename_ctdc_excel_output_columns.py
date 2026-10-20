#!/usr/bin/env python3
"""Rename CTDC input workbook output columns: TsvExcel→DbExcel, *_TSVData→*_DB, *_WebData→*_UI."""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
INPUT = ROOT / "InputFiles" / "CTDC"


def _rename_file_cell(value: str) -> str:
    if not value:
        return value
    s = str(value).strip()
    for old, new in (
        ("_TSVData.xlsx", "_DB.xlsx"),
        ("_TsvData.xlsx", "_DB.xlsx"),
        ("_WebData.xlsx", "_UI.xlsx"),
    ):
        s = s.replace(old, new)
    return s


def update_workbook(path: Path) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    tsv_col = headers.get("TsvExcel") or headers.get("tsvExcel") or headers.get("dbExcel")
    web_col = headers.get("WebExcel") or headers.get("UIExcel")
    if tsv_col:
        ws.cell(1, tsv_col, "DbExcel")
    for row in range(2, ws.max_row + 1):
        if tsv_col:
            v = ws.cell(row, tsv_col).value
            if v:
                ws.cell(row, tsv_col, _rename_file_cell(v))
        if web_col:
            v = ws.cell(row, web_col).value
            if v:
                ws.cell(row, web_col, _rename_file_cell(v))
    wb.save(path)
    wb.close()


def main():
    for path in sorted(INPUT.glob("TC*.xlsx")):
        update_workbook(path)
        print(path.name)


if __name__ == "__main__":
    main()
