#!/usr/bin/env python3
"""Generate Testcases/BENTO/arm-{a,d}/ tests and flat InputFiles/BENTO/TC*.xlsx from Katalon workbooks."""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
KATALON_INPUT = ROOT.parent / "Commons_Automation" / "InputFiles" / "Bento"
TEST_ROOT = ROOT / "Testcases" / "BENTO"
INPUT_ROOT = ROOT / "InputFiles" / "BENTO"

import sys

sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "Tools"))
sys.path.insert(0, str(ROOT / "Tools" / "cypher"))
from bento_arm_filters import (  # noqa: E402
    ARM_LABELS,
    ARM_SLUGS,
    ARMS,
    ArmLetter,
    BENTO_FACET_SPECS,
    ArmFacetChoice,
    BentoFacetSpec,
    choice_for_arm,
    workbook_relative_path,
)
from bento_cypher import build_queries_for_arm_test  # noqa: E402


def _copy_workbook(choice: ArmFacetChoice, dest: Path, arm: ArmLetter, spec: BentoFacetSpec) -> None:
    src = KATALON_INPUT / choice.katalon_workbook
    if not src.is_file():
        raise FileNotFoundError(src)
    shutil.copy2(src, dest)
    bundle = build_queries_for_arm_test(arm, spec, choice, KATALON_INPUT)
    wb = openpyxl.load_workbook(dest)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=2):
        tab = row[0].value
        if not tab:
            continue
        tab = str(tab)
        if tab == "CasesTab":
            row[1].value = bundle.cases_tab
            if len(row) > 2:
                row[2].value = bundle.stat_query
        elif tab == "SamplesTab":
            row[1].value = bundle.samples_tab
        elif tab == "FilesTab":
            row[1].value = bundle.files_tab
    wb.save(dest)
    wb.close()


def _filter_step_lines(arm: ArmLetter, spec: BentoFacetSpec, choice: ArmFacetChoice) -> list[str]:
    arm_line = f'    bento.select_filter("Filter/Arm", "{arm}")  # {ARM_LABELS[arm]}'
    if choice.age_min is not None:
        facet_line = (
            f"    bento.set_age_range({choice.age_min}, {choice.age_max})  "
            f"# {spec.facet_label} = {choice.value_label}"
        )
    else:
        facet_line = (
            f'    bento.select_filter("{spec.filter_group}", "{choice.filter_value}")  '
            f"# {spec.facet_label} = {choice.value_label}"
        )
    return [arm_line, facet_line]


def _test_py(arm: ArmLetter, spec: BentoFacetSpec, choice: ArmFacetChoice) -> str:
    slug = ARM_SLUGS[arm]
    wb = workbook_relative_path(arm, spec)
    test_name = f"test_TC{spec.tc:02d}_Bento_{slug.replace('-', '_')}_{choice.slug.replace('-', '_')}"
    doc = (
        f"Open Explore, Arm {arm} ({ARM_LABELS[arm]}) and "
        f"{spec.facet_label} {choice.value_label}, then compare stat bar and "
        f"Cases / Samples / Files tabs to expected queries."
    )
    filter_lines = "\n".join(_filter_step_lines(arm, spec, choice))
    return f'''"""{doc}"""

import pytest


@pytest.mark.bento
@pytest.mark.regression
def {test_name}(bento):
    bento.load_query_file("{wb}")
    bento.accept_banner()
    bento.open_cases()
{filter_lines}
    bento.read_statbar()
    bento.verify_tab("Cases")
    bento.verify_tab("Samples")
    bento.verify_tab("Files")
    bento.assert_no_failures()
'''


def _prune_old_workbooks() -> None:
    for path in INPUT_ROOT.glob("TC*_Bento_arm-*.xlsx"):
        path.unlink()


def _prune_old_tests(test_dir: Path) -> None:
    for path in test_dir.glob("test_TC*.py"):
        path.unlink()


def main() -> None:
    INPUT_ROOT.mkdir(parents=True, exist_ok=True)
    _prune_old_workbooks()

    for arm in ARMS:
        slug = ARM_SLUGS[arm]
        test_dir = TEST_ROOT / slug
        test_dir.mkdir(parents=True, exist_ok=True)
        _prune_old_tests(test_dir)
        for spec in BENTO_FACET_SPECS:
            choice = choice_for_arm(spec, arm)
            dest = INPUT_ROOT / workbook_relative_path(arm, spec)
            _copy_workbook(choice, dest, arm, spec)
            test_file = test_dir / f"test_TC{spec.tc:02d}_Bento_{slug}_{choice.slug.replace('-', '_')}.py"
            test_file.write_text(_test_py(arm, spec, choice))

    total = len(ARMS) * len(BENTO_FACET_SPECS)
    print(f"Generated {total} tests; workbooks in InputFiles/BENTO/ (flat, distinct checkbox per arm per facet)")


if __name__ == "__main__":
    main()
