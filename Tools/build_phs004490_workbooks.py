"""Generate InputFiles/C3DC TC01–TC04 workbooks for phs004490 (fast phs002827-style Cypher)."""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
IN = ROOT / "InputFiles" / "C3DC"
PHS = "phs004490"

STAT_TAIL = (
    "WITH DISTINCT p, s\n"
    "// Explore stat bar Files is 0 on QA for this cohort; Studies tab validates file totals.\n"
    "OPTIONAL MATCH (sm:sample)-[:of_sample]->(p)\n"
    "RETURN count(DISTINCT s) AS studies,\n"
    "       count(DISTINCT p) AS participants,\n"
    "       count(DISTINCT sm) AS samples,\n"
    "       0 AS files"
)

STAT_DIAG = (
    f"MATCH (d:diagnosis)-[:of_diagnosis]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)\n"
    f"WHERE s.dbgap_accession = '{PHS}' AND {{pred}}\n"
    + STAT_TAIL
)

STAT_SEX = (
    f"MATCH (p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)\n"
    f"WHERE s.dbgap_accession = '{PHS}' AND p.sex_at_birth = 'Male'\n"
    + STAT_TAIL
)

STAT_SAMPLE_SITE = (
    f"MATCH (sm:sample)-[:of_sample]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)\n"
    f"WHERE s.dbgap_accession = '{PHS}' AND 'chest wall' IN coalesce(sm.anatomic_site, [])\n"
    "WITH DISTINCT p, s, sm\n"
    "RETURN count(DISTINCT s) AS studies,\n"
    "       count(DISTINCT p) AS participants,\n"
    "       count(DISTINCT sm) AS samples,\n"
    "       0 AS files\n"
    "// Explore stat bar Files is 0 on QA for this cohort; Studies tab validates file totals."
)

STAT_CHEMO = (
    f"MATCH (t:treatment_chemotherapy)-[:of_treatment_chemotherapy]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)\n"
    f"WHERE s.dbgap_accession = '{PHS}' AND 'Chemotherapy' IN coalesce(t.treatment_type, [])\n"
    + STAT_TAIL
)


def clone_and_replace(src: str, dst: str, replacements: list[tuple[str, str]]) -> None:
    shutil.copy2(IN / src, IN / dst)
    wb = openpyxl.load_workbook(IN / dst)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val = cell.value
                for old, new in replacements:
                    val = val.replace(old, new)
                cell.value = val
    wb.save(IN / dst)
    wb.close()


def set_stat(dst: str, stat: str) -> None:
    wb = openpyxl.load_workbook(IN / dst)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "ParticipantsTab":
            row[2].value = stat
            break
    wb.save(IN / dst)
    wb.close()


def copy_studies_tab(from_dst: str, to_dsts: list[str], exists_prefix: str) -> None:
    wb = openpyxl.load_workbook(IN / from_dst, read_only=True, data_only=True)
    studies_q = None
    for row in wb.active.iter_rows(min_row=2, values_only=True):
        if row[0] == "StudiesTab":
            studies_q = row[1]
            break
    wb.close()
    if not studies_q:
        raise RuntimeError(f"No StudiesTab in {from_dst}")
    for dst in to_dsts:
        wb = openpyxl.load_workbook(IN / dst)
        ws = wb.active
        old_exists = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == "StudiesTab":
                old_exists = row[1].split("WITH s")[0]
                break
        if not old_exists:
            continue
        new_q = studies_q.replace(studies_q.split("WITH s")[0], exists_prefix)
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "StudiesTab":
                row[1].value = new_q
                break
        wb.save(IN / dst)
        wb.close()


def main() -> None:
    # TC01 — Diagnosis Basis Pathology (from phs004232 TC05 Clinical)
    clone_and_replace(
        "TC05_C3DC_phs004232_DiagnBasis-Clinical.xlsx",
        "TC01_C3DC_phs004490_DiagBasis-Pathology.xlsx",
        [("phs004232", PHS), ("'Clinical'", "'Pathology'"), ("Clinical", "Pathology")],
    )
    set_stat(
        "TC01_C3DC_phs004490_DiagBasis-Pathology.xlsx",
        STAT_DIAG.format(pred="'Pathology' IN coalesce(d.diagnosis_basis, [])"),
    )

    # TC02 — Sex at Birth Male (from phs004232 TC03 Race)
    clone_and_replace(
        "TC03_C3DC_phs004232_Race-Unknown.xlsx",
        "TC02_C3DC_phs004490_Sex-Male.xlsx",
        [
            ("phs004232", PHS),
            ("'Unknown' IN coalesce(p.race, [])", "p.sex_at_birth = 'Male'"),
            ("Race-Unknown", "Sex-Male"),
        ],
    )
    set_stat("TC02_C3DC_phs004490_Sex-Male.xlsx", STAT_SEX)

    # TC03 — Sample Anatomic Site chest wall (from phs004232 TC04 sample tumor)
    clone_and_replace(
        "TC04_C3DC_phs004232_SampleTumorStatus-NotReported.xlsx",
        "TC03_C3DC_phs004490_SampleSite-ChestWall.xlsx",
        [
            ("phs004232", PHS),
            ("sm.sample_tumor_status = 'Not Reported'", "'chest wall' IN coalesce(sm.anatomic_site, [])"),
            ("parent.sample_tumor_status = 'Not Reported'", "'chest wall' IN coalesce(parent.anatomic_site, [])"),
            ("SampleTumorStatus-NotReported", "SampleSite-ChestWall"),
        ],
    )
    set_stat("TC03_C3DC_phs004490_SampleSite-ChestWall.xlsx", STAT_SAMPLE_SITE)

    # TC04 — Treatment Type Chemotherapy (from phs002827 TC04 Radiation)
    clone_and_replace(
        "TC04_C3DC_phs002827_TrtmntType-Radiation.xlsx",
        "TC04_C3DC_phs004490_TrtType-Chemo.xlsx",
        [
            ("phs002827", PHS),
            ("treatment_radiation", "treatment_chemotherapy"),
            ("of_treatment_radiation", "of_treatment_chemotherapy"),
            ("treatment_radiation_id", "treatment_chemotherapy_id"),
            ("'Radiation'", "'Chemotherapy'"),
            ("TrtmntType-Radiation", "TrtType-Chemo"),
        ],
    )
    set_stat("TC04_C3DC_phs004490_TrtType-Chemo.xlsx", STAT_CHEMO)

    _apply_studies_tabs_from_phs002827(exists := {
        "TC01_C3DC_phs004490_DiagBasis-Pathology.xlsx": f"""MATCH (s:study)
WHERE EXISTS {{
  MATCH (d:diagnosis)-[:of_diagnosis]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s)
  WHERE s.dbgap_accession = '{PHS}' AND 'Pathology' IN coalesce(d.diagnosis_basis, [])
}}""",
        "TC02_C3DC_phs004490_Sex-Male.xlsx": f"""MATCH (s:study)
WHERE EXISTS {{
  MATCH (p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s)
  WHERE s.dbgap_accession = '{PHS}' AND p.sex_at_birth = 'Male'
}}""",
        "TC03_C3DC_phs004490_SampleSite-ChestWall.xlsx": f"""MATCH (s:study)
WHERE EXISTS {{
  MATCH (sm:sample)-[:of_sample]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s)
  WHERE s.dbgap_accession = '{PHS}' AND 'chest wall' IN coalesce(sm.anatomic_site, [])
}}""",
        "TC04_C3DC_phs004490_TrtType-Chemo.xlsx": f"""MATCH (s:study)
WHERE EXISTS {{
  MATCH (t:treatment_chemotherapy)-[:of_treatment_chemotherapy]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s)
  WHERE s.dbgap_accession = '{PHS}' AND 'Chemotherapy' IN coalesce(t.treatment_type, [])
}}""",
    })
    print("Wrote 4 workbooks for phs004490")


def _apply_studies_tabs_from_phs002827(exists: dict[str, str]) -> None:
    """phs002827 Studies shape (count DISTINCT d2, file_type tie-break) + Explore tie order for phs004490."""
    wb = openpyxl.load_workbook(IN / "TC04_C3DC_phs002827_TrtmntType-Radiation.xlsx", read_only=True, data_only=True)
    q827 = next(row[1] for row in wb.active.iter_rows(min_row=2, values_only=True) if row[0] == "StudiesTab")
    wb.close()
    body = "WITH s" + q827.split("WITH s", 1)[1]
    patches = [
        ("ORDER BY site_n DESC, site ASC", 
         "ORDER BY site_n DESC,\n"
         "  CASE site WHEN 'lymph node' THEN 1 WHEN 'upper lobe of lung' THEN 1 ELSE 0 END ASC,\n"
         "  CASE site WHEN 'diaphragm' THEN 1 WHEN 'mediastinum' THEN 2 WHEN 'liver' THEN 3 ELSE 4 END ASC,\n"
         "  site ASC"),
        ("ORDER BY type_n DESC\n", "ORDER BY type_n DESC, file_type ASC\n"),
    ]
    for dst, prefix in exists.items():
        q = prefix + body.replace("phs002827", PHS)
        for old, new in patches:
            q = q.replace(old, new)
        wb = openpyxl.load_workbook(IN / dst)
        ws = wb.active
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "StudiesTab":
                row[1].value = q
                break
        wb.save(IN / dst)
        wb.close()


if __name__ == "__main__":
    main()
