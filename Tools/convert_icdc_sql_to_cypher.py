#!/usr/bin/env python3
"""
Replace SQL in InputFiles/ICDC/TC*.xlsx with Neo4j Cypher (StatQuery + 4 tab rows).

    python Tools/convert_icdc_sql_to_cypher.py
    python Tools/convert_icdc_sql_to_cypher.py --dry-run

Renames output workbook columns to C3DC-style *_DbData.xlsx / *_WebData.xlsx (no TSV).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Tools" / "cypher"))

import icdc_cypher as ic  # noqa: E402

INPUT = ROOT / "InputFiles" / "ICDC"
HEADER = ["TabName", "TabQuery", "StatQuery", "TsvExcel", "WebExcel"]


def _rows_by_tab(ws) -> dict[str, list]:
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=False):
        key = row[0].value
        if not key:
            continue
        out[str(key).strip()] = row
    return out


def _output_names(stem: str) -> tuple[str, str]:
    return f"{stem}_DbData.xlsx", f"{stem}_WebData.xlsx"


def convert_workbook(path: Path, dry_run: bool) -> bool:
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    if ws.max_row < 2:
        wb.close()
        return False
    rows = _rows_by_tab(ws)
    cases_row = rows.get("CasesTab")
    if not cases_row or not cases_row[1].value:
        wb.close()
        return False
    cases_sql = str(cases_row[1].value)
    try:
        cy = ic.all_queries_from_cases_sql(cases_sql)
    except ValueError as exc:
        print(f"skip {path.name}: {exc}")
        wb.close()
        return False
    db_name, web_name = _output_names(path.stem)
    if dry_run:
        print(f"{path.name}: would write Cypher + {db_name}, {web_name}")
        wb.close()
        return True
    if ws.max_row == 1 or [c.value for c in ws[1]][:5] != HEADER:
        ws.delete_rows(1, ws.max_row)
        ws.append(HEADER)
        rows = {}
    for tab, cypher in cy.items():
        if tab == "StatQuery":
            cases_row[2].value = cypher
            continue
        if tab in rows:
            rows[tab][1].value = cypher
        else:
            ws.append([tab, cypher, None, None, None])
            rows[tab] = ws[ws.max_row]
    cases_row[3].value = db_name
    cases_row[4].value = web_name
    wb.save(path)
    wb.close()
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--input", type=Path, default=INPUT)
    args = ap.parse_args()
    n = 0
    for path in sorted(args.input.glob("TC*.xlsx")):
        if convert_workbook(path, args.dry_run):
            n += 1
            if not args.dry_run:
                print(f"updated {path.name}")
    print(f"done: {n} workbook(s)")


if __name__ == "__main__":
    main()
