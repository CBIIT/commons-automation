#!/usr/bin/env python3
"""Write Neo4j Cypher into InputFiles/BENTO/TC*_Bento_arm-*.xlsx from Katalon sources + facet specs."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
KATALON_INPUT = ROOT.parent / "Commons_Automation" / "InputFiles" / "Bento"
INPUT_ROOT = ROOT / "InputFiles" / "BENTO"

sys.path.insert(0, str(ROOT))
from Commons.bento_arm_filters import (  # noqa: E402
    ARMS,
    BENTO_FACET_SPECS,
    ArmLetter,
    choice_for_arm,
    workbook_file_name,
)
from Commons.bento_cypher_queries import BentoQueryBundle, build_queries_for_arm_test  # noqa: E402


def _write_workbook(path: Path, bundle: BentoQueryBundle) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=2):
        tab_cell = row[0]
        if not tab_cell.value:
            continue
        tab = str(tab_cell.value)
        query_cell = row[1]
        stat_cell = row[2] if len(row) > 2 else None
        if tab == "CasesTab":
            query_cell.value = bundle.cases_tab
            if stat_cell is not None:
                stat_cell.value = bundle.stat_query
        elif tab == "SamplesTab":
            query_cell.value = bundle.samples_tab
        elif tab == "FilesTab":
            query_cell.value = bundle.files_tab
    wb.save(path)
    wb.close()


def main() -> None:
    count = 0
    for arm in ARMS:
        for spec in BENTO_FACET_SPECS:
            choice = choice_for_arm(spec, arm)
            dest = INPUT_ROOT / workbook_file_name(arm, spec)
            if not dest.is_file():
                raise FileNotFoundError(dest)
            bundle = build_queries_for_arm_test(arm, spec, choice, KATALON_INPUT)
            _write_workbook(dest, bundle)
            count += 1
    print(f"Updated Cypher in {count} workbooks under {INPUT_ROOT}")


if __name__ == "__main__":
    main()
