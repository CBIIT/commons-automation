"""Write GC Excel test case workbooks (Steps sheet)."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

HEADER = ["Input Excel", "Action", "Page", "Params", "Filter"]

GC_EXCEL_ACTIONS = [
    "accept_banner",
    "open_data",
    "select_filter",
    "read_statbar",
    "verify_tab",
    "verify_all_tabs",
    "assert_no_failures",
    "click_and_wait",
    "js_click_and_wait",
]

_HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")


def write_gc_excel_case(
    out_path: Path,
    *,
    input_excel: str,
    steps: list[tuple[str, str, str, str]],
) -> None:
    """steps: (action, page, params, filter) — skip load_query_file (row 2 holds Input Excel)."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Steps"
    ws.append(HEADER)
    ws.append([input_excel, None, None, None, None])
    for action, page, params, flt in steps:
        ws.append([None, action, page, params, flt])

    lists = wb.create_sheet("_Lists")
    lists.sheet_state = "hidden"
    for i, act in enumerate(GC_EXCEL_ACTIONS, start=1):
        lists.cell(row=i, column=1, value=act)
    wb.defined_names.add(DefinedName("GCExcelActions", attr_text="_Lists!$A$1:$A$50"))

    dv = DataValidation(type="list", formula1="=GCExcelActions", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B3:B500")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    wb.save(out_path)
