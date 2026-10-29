"""One-off: generate InputFiles/C3DC TC01–TC05 workbooks for phs004232 from existing templates."""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
IN = ROOT / "InputFiles" / "C3DC"
PHS = "phs004232"

STAT_UNION_FROM_PARTICIPANTS = """
MATCH (p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)
WHERE s.dbgap_accession = '{phs}'
  AND {cohort_predicate}
WITH collect(DISTINCT s) AS study_nodes, collect(DISTINCT p) AS part_nodes
CALL {{
  WITH part_nodes
  UNWIND part_nodes AS p
  MATCH (sm:sample)-[:of_sample]->(p)
  RETURN sm
  UNION
  WITH part_nodes
  UNWIND part_nodes AS p
  MATCH (parent:sample)-[:of_sample]->(p)
  MATCH (sm:sample)-[:of_sample]->(:cell_line)-[:of_cell_line]->(parent)
  RETURN sm
  UNION
  WITH part_nodes
  UNWIND part_nodes AS p
  MATCH (parent:sample)-[:of_sample]->(p)
  MATCH (sm:sample)-[:of_sample]->(:pdx)-[:of_pdx]->(parent)
  RETURN sm
}}
WITH study_nodes, part_nodes, collect(DISTINCT sm) AS sample_nodes
RETURN size(study_nodes) AS studies,
       size(part_nodes) AS participants,
       size(sample_nodes) AS samples,
       0 AS files
// Explore stat bar Files is 0 on QA for this cohort; Studies tab validates file totals.
""".strip()

STAT_SAMPLE_TUMOR = """
CALL {{
  MATCH (sm:sample)-[:of_sample]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)
  WHERE s.dbgap_accession = '{phs}' AND sm.sample_tumor_status = 'Not Reported'
  RETURN sm, p, s
  UNION
  MATCH (parent:sample)-[:of_sample]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)
  WHERE s.dbgap_accession = '{phs}' AND parent.sample_tumor_status = 'Not Reported'
  MATCH (sm:sample)-[:of_sample]->(:cell_line)-[:of_cell_line]->(parent)
  RETURN sm, p, s
  UNION
  MATCH (parent:sample)-[:of_sample]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)
  WHERE s.dbgap_accession = '{phs}' AND parent.sample_tumor_status = 'Not Reported'
  MATCH (sm:sample)-[:of_sample]->(:pdx)-[:of_pdx]->(parent)
  RETURN sm, p, s
}}
WITH DISTINCT s, p, sm
RETURN count(DISTINCT s) AS studies,
       count(DISTINCT p) AS participants,
       count(DISTINCT sm) AS samples,
       0 AS files
// Explore stat bar Files is 0 on QA for this cohort; Studies tab validates file totals.
""".strip()


def clone_and_replace(src: str, dst: str, replacements: list[tuple[str, str]]) -> None:
    src_path = IN / src
    dst_path = IN / dst
    shutil.copy2(src_path, dst_path)
    wb = openpyxl.load_workbook(dst_path)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val = cell.value
                for old, new in replacements:
                    val = val.replace(old, new)
                cell.value = val
    wb.save(dst_path)
    wb.close()


def set_stat(dst: str, stat: str) -> None:
    wb = openpyxl.load_workbook(IN / dst)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "ParticipantsTab":
            row[2].value = stat
            break
    wb.save(IN / dst)
    wb.close()


def main() -> None:
    # TC01 — Study Phase Completed (from phs003162 TC01)
    clone_and_replace(
        "TC01_C3DC_phs003162_StudyPhase-Completed.xlsx",
        "TC01_C3DC_phs004232_StudyPhase-Completed.xlsx",
        [("phs003162", PHS), ("2 AS files", "0 AS files")],
    )

    # TC02 — Diagnosis Anatomic Site bone element
    clone_and_replace(
        "TC04_C3DC_phs002430_AnatoSite-C119NasopharynxNOS.xlsx",
        "TC02_C3DC_phs004232_AnatoSite-BoneElement.xlsx",
        [("phs002430", PHS), ("'nasopharynx'", "'bone element'")],
    )
    set_stat(
        "TC02_C3DC_phs004232_AnatoSite-BoneElement.xlsx",
        STAT_UNION_FROM_PARTICIPANTS.format(
            phs=PHS,
            cohort_predicate=(
                "EXISTS { MATCH (d:diagnosis)-[:of_diagnosis]->(p) "
                "WHERE 'bone element' IN coalesce(d.anatomic_site, []) }"
            ),
        ),
    )

    # TC03 — Race Unknown (from phs002827 race TC03)
    clone_and_replace(
        "TC03_C3DC_phs002827_Race-MidEastOrNorthAfrican.xlsx",
        "TC03_C3DC_phs004232_Race-Unknown.xlsx",
        [
            ("phs002827", PHS),
            ("'Middle Eastern or North African'", "'Unknown'"),
        ],
    )
    set_stat(
        "TC03_C3DC_phs004232_Race-Unknown.xlsx",
        STAT_UNION_FROM_PARTICIPANTS.format(
            phs=PHS,
            cohort_predicate="'Unknown' IN coalesce(p.race, [])",
        ),
    )

    # TC04 — Sample Tumor Status Not Reported (from phs003162 tumor extent Local)
    clone_and_replace(
        "TC04_C3DC_phs003162_TumorExtent-Local.xlsx",
        "TC04_C3DC_phs004232_SampleTumorStatus-NotReported.xlsx",
        [
            ("phs003162", PHS),
            ("sm.tumor_spatial_extent = 'Local'", "sm.sample_tumor_status = 'Not Reported'"),
            ("parent.tumor_spatial_extent = 'Local'", "parent.sample_tumor_status = 'Not Reported'"),
        ],
    )
    set_stat(
        "TC04_C3DC_phs004232_SampleTumorStatus-NotReported.xlsx",
        STAT_SAMPLE_TUMOR.format(phs=PHS),
    )

    # TC05 — Diagnosis Basis Clinical
    clone_and_replace(
        "TC07_C3DC_phs002430_DiagnBasis-Clinical.xlsx",
        "TC05_C3DC_phs004232_DiagnBasis-Clinical.xlsx",
        [("phs002430", PHS)],
    )

    _patch_samples_union_diagnosis(
        "TC02_C3DC_phs004232_AnatoSite-BoneElement.xlsx",
        "'bone element' IN coalesce(d.anatomic_site, [])",
    )
    _patch_samples_union_participant(
        "TC03_C3DC_phs004232_Race-Unknown.xlsx",
        "'Unknown' IN coalesce(p.race, [])",
    )
    _patch_samples_union_diagnosis(
        "TC05_C3DC_phs004232_DiagnBasis-Clinical.xlsx",
        "'Clinical' IN coalesce(d.diagnosis_basis, [])",
    )

    print("Wrote 5 workbooks under InputFiles/C3DC/")


_SAMPLES_RETURN_TAIL = (
    "WITH DISTINCT sm, p, s\n"
    "RETURN p.participant_id AS `Participant ID`,\n"
    "       sm.sample_id AS `Sample ID`,\n"
    "       s.study_id AS `Study ID`,\n"
    "       sm.tumor_spatial_extent AS `Tumor Spatial Extent`,\n"
    "       reduce(text = '', site IN coalesce(sm.anatomic_site, []) | CASE WHEN text = '' THEN site ELSE text + ';' + site END) AS `Anatomic Site`,\n"
    "       CASE\n"
    "         WHEN sm.participant_age_at_collection IS NULL OR sm.participant_age_at_collection = -999 THEN 'Not Reported'\n"
    "         ELSE toString(sm.participant_age_at_collection)\n"
    "       END AS `Participant Age at Collection (days)`,\n"
    "       sm.sample_tumor_status AS `Sample Tumor Status`\n"
    "ORDER BY sm.sample_id\n"
    "LIMIT 100;\n"
)


def _union_call(*branches: str) -> str:
    return "CALL {\n" + "\n  UNION\n".join(branches) + "\n}\n" + _SAMPLES_RETURN_TAIL


def _sample_branches_from_anchor(anchor: str) -> str:
    direct = (
        anchor
        + "  MATCH (sm:sample)-[:of_sample]->(p)\n"
        + "  RETURN sm, p, s"
    )
    cell_line = (
        anchor
        + "  MATCH (parent:sample)-[:of_sample]->(p)\n"
        + "  MATCH (sm:sample)-[:of_sample]->(:cell_line)-[:of_cell_line]->(parent)\n"
        + "  RETURN sm, p, s"
    )
    pdx = (
        anchor
        + "  MATCH (parent:sample)-[:of_sample]->(p)\n"
        + "  MATCH (sm:sample)-[:of_sample]->(:pdx)-[:of_pdx]->(parent)\n"
        + "  RETURN sm, p, s"
    )
    return _union_call(direct, cell_line, pdx)


def _patch_samples_union_participant(dst: str, cohort_line: str) -> None:
    anchor = (
        f"  MATCH (p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)\n"
        f"  WHERE s.dbgap_accession = '{PHS}' AND {cohort_line}\n"
        "  WITH DISTINCT p, s\n"
    )
    _write_samples_tab(dst, _sample_branches_from_anchor(anchor))


def _patch_samples_union_diagnosis(dst: str, diagnosis_predicate: str) -> None:
    # Cohort-first on diagnosis (see phs003162 TC02 Gliomas) — avoids slow nested EXISTS.
    anchor = (
        "  MATCH (d:diagnosis)-[:of_diagnosis]->(p:participant)-[:of_participant]->(:consent_group)-[:of_consent_group]->(s:study)\n"
        f"  WHERE s.dbgap_accession = '{PHS}' AND {diagnosis_predicate}\n"
        "  WITH DISTINCT p, s\n"
    )
    _write_samples_tab(dst, _sample_branches_from_anchor(anchor))


def _write_samples_tab(dst: str, query: str) -> None:
    wb = openpyxl.load_workbook(IN / dst)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "SamplesTab":
            row[1].value = query
            break
    wb.save(IN / dst)
    wb.close()


if __name__ == "__main__":
    main()
